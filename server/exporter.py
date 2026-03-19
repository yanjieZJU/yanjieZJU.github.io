import os
import base64
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import json
import httpx
from openai import OpenAI

import dashscope
from dashscope import MultiModalConversation

dashscope.base_http_api_url = 'https://dashscope.aliyuncs.com/api/v1'
api_key = os.environ.get("DASHSCOPE_API_KEY")
http_client = httpx.Client(proxy=None, trust_env=False)
client = OpenAI(
    http_client=http_client,
    api_key=os.environ.get("OPENAI_API_KEY"),
    base_url=os.environ.get("OPENAI_BASE_URL")
)

def clean_plan_for_prompt(plan):
    """
    Remove technical_methods, references, and other verbose fields
    that don't help with image generation.
    """
    if isinstance(plan, dict):
        modules = list(plan.values())
    else:
        modules = plan

    cleaned = []
    for mod in modules:
        cleaned_mod = {
            "name": mod.get("name", ""),
            "brief_description": mod.get("brief_description", ""),
            "role": mod.get("role", ""),
        }
        # Include simplified sub_nodes if present
        if mod.get("sub_nodes"):
            cleaned_mod["sub_nodes"] = [
                {"name": sn.get("name", ""), "brief_description": sn.get("brief_description", "")}
                for sn in mod["sub_nodes"]
            ]
        # Include simplified ports for connectivity info
        if mod.get("ports"):
            ports = mod["ports"]
            if ports.get("outgoing"):
                cleaned_mod["connections_to"] = [
                    edge.get("target_mod", "") for edge in ports["outgoing"]
                ]
        cleaned.append(cleaned_mod)

    return cleaned

def generate_image_prompt(plan_json):
    """
    Use GPT to convert the architecture JSON into a structured image generation prompt
    following the S-C-S-S framework.
    """
    cleaned = clean_plan_for_prompt(plan_json)

    system_prompt = """你是一位专业的系统架构可视化专家。你的任务是将系统架构方案转化为适合AI绘图工具的提示词。

    提示词应遵循 S-C-S-S 框架：
    - Subject (主体): 图中要呈现的核心对象和模块
    - Composition (构图): 整体布局、信息块之间的空间关系和层次
    - Structure (结构细节): 每个模块的内部元素、数据流箭头走向、连接关系
    - Style (风格渲染): 配色方案、图形风格（扁平/立体）、边框样式、质感

    输出格式要求：
    1. 简洁明确，避免冗余描述
    2. 突出模块间的数据流和层次关系
    3. 使用专业的架构图术语（如"模块"、"数据流"、"接口"等）
    4. 配色建议应与模块角色（input/processing/output/data）相关联
    5. 整体风格应为现代、清晰、专业的技术架构图

    请直接输出可用于绘图的提示词，不要包含解释性文字。"""

    user_message = f"""下面是一个系统架构方案，请将它整理成可直接给绘图工具使用的 prompt：

    {json.dumps(cleaned, ensure_ascii=False, indent=2)}

    请按照 S-C-S-S 框架生成绘图提示词。"""

    try:
        response = client.chat.completions.create(
            model="gpt-5.2-chat-latest",
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_message}
            ],
            temperature=0.2,
        )
        prompt = response.choices[0].message.content.strip()
        print(f"✅ Generated image prompt ({len(prompt)} chars)")
        return prompt
    except Exception as e:
        print(f"⚠️  Failed to generate image prompt: {e}")
        # Fallback: simple description
        module_names = [m.get("name", "") for m in cleaned]
        return f"A professional system architecture diagram showing modules: {', '.join(module_names)}. Clean modern style with clear data flow arrows."

def generate_diagram_image(plan_json):
    """
    Generate a high-quality architecture diagram using Qwen image generation model.
    Returns base64-encoded PNG image data (without data URI prefix).
    """
    if not api_key:
        print("⚠️  DASHSCOPE_API_KEY not set, skipping diagram generation")
        return None

    print("🎨 Generating architecture diagram with Qwen...")

    # Step 1: Generate structured prompt
    image_prompt = generate_image_prompt(plan_json)

    # Step 2: Call Qwen image generation
    messages = [{
        "role": "user",
        "content": [
            {"text": image_prompt}
        ]
    }]

    try:
        response = MultiModalConversation.call(
            api_key=api_key,
            model="qwen-image-2.0-pro",
            messages=messages,
            result_format='message',
            stream=False,
            watermark=False,
            prompt_extend=True,
            size='1536*1152'  # 4:3 ratio, high resolution
        )

        if response.status_code == 200:
            # Extract image URL from response
            output = response.output
            if output and 'choices' in output:
                choice = output['choices'][0]
                if 'message' in choice and 'content' in choice['message']:
                    content = choice['message']['content']
                    if isinstance(content, list):
                        for item in content:
                            if isinstance(item, dict) and 'image' in item:
                                image_url = item['image']
                                print(f"✅ Generated diagram image: {image_url[:60]}...")

                                # Download the image and convert to base64
                                img_response = httpx.get(image_url, timeout=30)
                                if img_response.status_code == 200:
                                    img_b64 = base64.b64encode(img_response.content).decode('utf-8')
                                    return img_b64

        print(f"⚠️  Unexpected response format: {response}")
        return None

    except Exception as e:
        print(f"❌ Failed to generate diagram: {e}")
        return None

def add_table_row(table, key, value, bold_key=True):
    row = table.add_row().cells
    
    # Key cell
    key_run = row[0].paragraphs[0].add_run(key)
    if bold_key:
        key_run.bold = True
    row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Value cell
    row[1].text = str(value) if value else "N/A"
    
def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = widths[idx]

def export_to_docx(image_b64: str, node_details: dict, overall_description: str = None) -> BytesIO:
    """
    Generate a Word document containing the architecture flowchart image,
    overall description, and detailed tables for each module.
    """
    doc = Document()

    # Set default font for Chinese characters support
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Microsoft YaHei'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    # Title
    title = doc.add_heading("System Architecture Plan", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- AI-generated diagram (placed at the very top) ---
    # print("🖼️  Generating AI diagram for export...")
    # ai_image_b64 = generate_diagram_image(node_details)
    # if ai_image_b64:
    #     doc.add_heading("1. AI-Generated Architecture Diagram", level=1)
    #     try:
    #         ai_image_data = base64.b64decode(ai_image_b64)
    #         ai_image_stream = BytesIO(ai_image_data)
    #         doc.add_picture(ai_image_stream, width=Inches(6.0))
    #         doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    #     except Exception as e:
    #         doc.add_paragraph(f"[Error rendering AI diagram: {str(e)}]")
    #     doc.add_paragraph("")
    # else:
    #     doc.add_heading("1. AI-Generated Architecture Diagram", level=1)
    #     doc.add_paragraph("[AI diagram generation skipped or failed]")

    # Overall Description
    if overall_description:
        doc.add_heading("Overall Description", level=1)
        doc.add_paragraph(overall_description)
        doc.add_paragraph("")

    # Add Flowchart Image (from frontend canvas)
    doc.add_heading("2. Architecture Flowchart", level=1)

    if image_b64:
        if ',' in image_b64:
            image_b64 = image_b64.split(',')[1]
        try:
            image_data = base64.b64decode(image_b64)
            image_stream = BytesIO(image_data)
            doc.add_picture(image_stream, width=Inches(6.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception as e:
            doc.add_paragraph(f"[Error rendering image: {str(e)}]")
    else:
        doc.add_paragraph("[No flowchart image provided]")

    doc.add_page_break()

    # Add Node Details
    doc.add_heading("3. Module Details", level=1)

    if isinstance(node_details, dict):
        modules = list(node_details.values())
    else:
        modules = node_details

    if not modules:
        doc.add_paragraph("No detailed module information available.")

    for idx, mod in enumerate(modules):
        mod_name = mod.get("name", f"Module {idx+1}")
        doc.add_heading(f"3.{idx+1} {mod_name}", level=2)

        table = doc.add_table(rows=0, cols=2)
        table.style = 'Table Grid'

        add_table_row(table, "Brief Description", mod.get("brief_description", ""))
        add_table_row(table, "Function", mod.get("function", ""))
        add_table_row(table, "Technical Methods", mod.get("technical_methods", ""))

        if "role" in mod:
            add_table_row(table, "Role", mod.get("role"))
        if "references" in mod:
            add_table_row(table, "References", mod.get("references"))

        set_col_widths(table, [Inches(1.5), Inches(4.5)])
        doc.add_paragraph("")

        sub_nodes = mod.get("sub_nodes", [])
        if sub_nodes:
            doc.add_heading(f"Sub-nodes for {mod_name}", level=3)
            for sub in sub_nodes:
                sn_name = sub.get("name", "Unknown Sub-node")
                p = doc.add_paragraph()
                p.add_run(f"• {sn_name}").bold = True

                sn_table = doc.add_table(rows=0, cols=2)
                sn_table.style = 'Light Shading Accent 1'

                add_table_row(sn_table, "Function", sub.get("function", ""))
                add_table_row(sn_table, "Technical Methods", sub.get("technical_methods", ""))
                set_col_widths(sn_table, [Inches(1.5), Inches(4.5)])
                doc.add_paragraph("")

    doc_stream = BytesIO()
    doc.save(doc_stream)
    doc_stream.seek(0)
    return doc_stream
